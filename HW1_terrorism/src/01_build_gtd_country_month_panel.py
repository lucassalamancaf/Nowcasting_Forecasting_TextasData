"""Build a complete GTD country-month panel.

This script loads an event-level Global Terrorism Database (GTD) file from
data/raw/, cleans the event dates, aggregates attacks to country-month counts,
and creates a complete country-month panel with zero rows for months without
attacks.

It intentionally does not define onset or forecasting labels.
"""

from __future__ import annotations

import argparse
import io
import textwrap
import zipfile
from pathlib import Path
from typing import Iterable

import pandas as pd


EXPECTED_COLUMNS = [
    "eventid",
    "iyear",
    "imonth",
    "iday",
    "country",
    "country_txt",
    "region",
    "region_txt",
    "provstate",
    "city",
    "latitude",
    "longitude",
    "doubtterr",
    "multiple",
    "success",
    "suicide",
    "attacktype1",
    "attacktype1_txt",
    "targtype1",
    "targtype1_txt",
    "weaptype1",
    "weaptype1_txt",
    "nkill",
    "nwound",
    "gname",
]

CORE_COLUMNS = [
    "eventid",
    "iyear",
    "imonth",
    "country",
    "country_txt",
    "region",
    "region_txt",
]

RAW_SUFFIXES = {".xlsx", ".xls", ".csv", ".zip"}


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Build a complete country-month panel from GTD event data."
    )
    parser.add_argument(
        "--start",
        default="2010-01",
        help="Start month for the panel, e.g. 2010-01 or 2010-01-01.",
    )
    parser.add_argument(
        "--end",
        default="2020-12",
        help="End month for the panel, e.g. 2020-12 or 2020-12-01.",
    )
    parser.add_argument(
        "--raw-dir",
        default="data/raw",
        type=Path,
        help="Directory containing the raw GTD file.",
    )
    parser.add_argument(
        "--interim-dir",
        default="data/interim",
        type=Path,
        help="Directory for cleaned event-level outputs.",
    )
    parser.add_argument(
        "--processed-dir",
        default="data/processed",
        type=Path,
        help="Directory for final panel outputs.",
    )
    parser.add_argument(
        "--diagnostics-dir",
        default="outputs/diagnostics",
        type=Path,
        help="Directory for diagnostics output.",
    )
    parser.add_argument(
        "--figures-dir",
        default="outputs/figures",
        type=Path,
        help="Directory for diagnostic plots.",
    )
    return parser.parse_args()


def month_start(value: str) -> pd.Timestamp:
    """Convert a YYYY-MM or date-like value to a month-start Timestamp."""
    ts = pd.to_datetime(value)
    return pd.Timestamp(year=ts.year, month=ts.month, day=1)


def ensure_output_dirs(paths: Iterable[Path]) -> None:
    """Create output directories if they do not exist."""
    for path in paths:
        path.mkdir(parents=True, exist_ok=True)


def find_gtd_file(raw_dir: Path) -> Path:
    """Find a likely raw GTD file in data/raw/."""
    if not raw_dir.exists():
        raise FileNotFoundError(
            f"Raw data directory does not exist: {raw_dir}\n"
            "Create it and place the GTD .xlsx, .xls, .csv, or .zip file there."
        )

    candidates = [
        path
        for path in raw_dir.iterdir()
        if path.is_file() and path.suffix.lower() in RAW_SUFFIXES
    ]
    if not candidates:
        raise FileNotFoundError(
            "No GTD file found in data/raw/.\n\n"
            "Download the Global Terrorism Database from START after accepting "
            "its terms of use, then place the .xlsx, .xls, .csv, or .zip file "
            f"in: {raw_dir.resolve()}"
        )

    def sort_key(path: Path) -> tuple[int, int, str]:
        name = path.name.lower()
        looks_like_gtd = any(token in name for token in ["gtd", "globalterrorism"])
        return (0 if looks_like_gtd else 1, -path.stat().st_size, name)

    candidates.sort(key=sort_key)
    return candidates[0]


def standardize_column_names(df: pd.DataFrame) -> pd.DataFrame:
    """Return a copy with stripped, lower-case column names."""
    df = df.copy()
    df.columns = [str(column).strip().lower() for column in df.columns]
    return df


def is_expected_column(column: object) -> bool:
    """Return True if a raw column name is needed for this panel build."""
    return str(column).strip().lower() in EXPECTED_COLUMNS


def read_csv_robust(source: Path | io.BytesIO) -> pd.DataFrame:
    """Read a CSV, trying common encodings used by public datasets."""
    encodings = ["utf-8", "utf-8-sig", "latin1", "ISO-8859-1"]
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            if hasattr(source, "seek"):
                source.seek(0)
            return pd.read_csv(
                source,
                low_memory=False,
                encoding=encoding,
                usecols=is_expected_column,
            )
        except UnicodeDecodeError as exc:
            last_error = exc
    raise UnicodeDecodeError(
        "utf-8",
        b"",
        0,
        1,
        f"Could not read CSV with common encodings. Last error: {last_error}",
    )


def read_xlsx_streaming(source: Path | io.BytesIO) -> pd.DataFrame:
    """Read needed columns from a large .xlsx workbook with openpyxl streaming."""
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook.active

    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_header = next(rows)
    except StopIteration:
        workbook.close()
        return pd.DataFrame()

    normalized_header = [str(column).strip().lower() for column in raw_header]
    selected = [
        (index, column)
        for index, column in enumerate(normalized_header)
        if column in EXPECTED_COLUMNS
    ]

    records = []
    selected_indices = [index for index, _ in selected]
    selected_columns = [column for _, column in selected]
    for row in rows:
        records.append(
            [row[index] if index < len(row) else None for index in selected_indices]
        )

    workbook.close()
    return pd.DataFrame(records, columns=selected_columns)


def choose_zip_member(zip_file: zipfile.ZipFile) -> str:
    """Choose a GTD-like data file from a zip archive."""
    members = [
        name
        for name in zip_file.namelist()
        if not name.endswith("/")
        and Path(name).suffix.lower() in {".xlsx", ".xls", ".csv"}
        and "__macosx" not in name.lower()
    ]
    if not members:
        raise ValueError("Zip archive does not contain a .xlsx, .xls, or .csv file.")

    def sort_key(name: str) -> tuple[int, str]:
        lowered = name.lower()
        looks_like_gtd = any(token in lowered for token in ["gtd", "globalterrorism"])
        return (0 if looks_like_gtd else 1, lowered)

    members.sort(key=sort_key)
    return members[0]


def load_gtd_file(path: Path) -> pd.DataFrame:
    """Load a GTD file from Excel, CSV, or zip format."""
    suffix = path.suffix.lower()
    print(f"Loading GTD file: {path}")
    print("Reading only the GTD columns needed for this panel build.")
    if suffix == ".csv":
        df = read_csv_robust(path)
    elif suffix == ".xlsx":
        df = read_xlsx_streaming(path)
    elif suffix == ".xls":
        df = pd.read_excel(path, usecols=is_expected_column)
    elif suffix == ".zip":
        with zipfile.ZipFile(path) as archive:
            member = choose_zip_member(archive)
            print(f"Reading zipped GTD member: {member}")
            data = io.BytesIO(archive.read(member))
            member_suffix = Path(member).suffix.lower()
            if member_suffix == ".csv":
                df = read_csv_robust(data)
            elif member_suffix == ".xlsx":
                df = read_xlsx_streaming(data)
            else:
                df = pd.read_excel(data, usecols=is_expected_column)
    else:
        raise ValueError(f"Unsupported file format: {path.suffix}")

    return standardize_column_names(df)


def validate_columns(df: pd.DataFrame) -> tuple[list[str], list[str]]:
    """Check expected and core GTD variables."""
    available = [column for column in EXPECTED_COLUMNS if column in df.columns]
    missing = [column for column in EXPECTED_COLUMNS if column not in df.columns]
    missing_core = [column for column in CORE_COLUMNS if column not in df.columns]

    if missing:
        print("Missing expected GTD variables:", ", ".join(missing))
    else:
        print("All expected GTD variables are available.")

    if missing_core:
        raise ValueError(
            "Raw GTD file is missing required variables: "
            + ", ".join(missing_core)
            + "\nCannot build the country-month panel without these columns."
        )
    return available, missing


def clean_event_data(
    raw_df: pd.DataFrame, available_columns: list[str]
) -> tuple[pd.DataFrame, dict[str, int]]:
    """Keep needed columns and create a month-start date for each event."""
    events = raw_df[available_columns].copy()
    raw_rows = len(events)

    events["iyear"] = pd.to_numeric(events["iyear"], errors="coerce")
    events["imonth"] = pd.to_numeric(events["imonth"], errors="coerce")

    missing_or_unknown_month = events["imonth"].isna() | (events["imonth"] == 0)
    dropped_unknown_month = int(missing_or_unknown_month.sum())
    events = events.loc[~missing_or_unknown_month].copy()

    invalid_month = ~events["imonth"].between(1, 12)
    dropped_invalid_month = int(invalid_month.sum())
    events = events.loc[~invalid_month].copy()

    events["month"] = pd.to_datetime(
        {
            "year": events["iyear"].astype("Int64"),
            "month": events["imonth"].astype("Int64"),
            "day": 1,
        },
        errors="coerce",
    )

    invalid_date = events["month"].isna()
    dropped_invalid_date = int(invalid_date.sum())
    events = events.loc[~invalid_date].copy()

    events["iyear"] = events["iyear"].astype("Int64")
    events["imonth"] = events["imonth"].astype("Int64")

    diagnostics = {
        "raw_rows_loaded": raw_rows,
        "dropped_unknown_month": dropped_unknown_month,
        "dropped_invalid_month": dropped_invalid_month,
        "dropped_invalid_date": dropped_invalid_date,
        "cleaned_event_rows": len(events),
    }
    return events, diagnostics


def numeric_series(df: pd.DataFrame, column: str, default: float = 0) -> pd.Series:
    """Return a numeric series, or a default-valued series if the column is missing."""
    if column in df.columns:
        return pd.to_numeric(df[column], errors="coerce")
    return pd.Series(default, index=df.index, dtype="float64")


def add_event_metrics(events: pd.DataFrame) -> pd.DataFrame:
    """Create per-event metric columns used for country-month aggregation."""
    events = events.copy()

    if "doubtterr" in events.columns:
        doubtterr = numeric_series(events, "doubtterr")
        strict = doubtterr.ne(1) | doubtterr.isna()
    else:
        strict = pd.Series(True, index=events.index)

    success = numeric_series(events, "success")
    suicide = numeric_series(events, "suicide")
    nkill = numeric_series(events, "nkill").fillna(0)
    nwound = numeric_series(events, "nwound").fillna(0)

    events["_strict_event"] = strict.astype(int)
    events["_n_attacks_all"] = 1
    events["_n_attacks_strict"] = events["_strict_event"]
    events["_n_success_all"] = success.eq(1).astype(int)
    events["_n_success_strict"] = events["_n_success_all"] * events["_strict_event"]
    events["_n_suicide_all"] = suicide.eq(1).astype(int)
    events["_n_suicide_strict"] = events["_n_suicide_all"] * events["_strict_event"]
    events["_fatalities_all"] = nkill
    events["_fatalities_strict"] = nkill * events["_strict_event"]
    events["_injuries_all"] = nwound
    events["_injuries_strict"] = nwound * events["_strict_event"]
    events["_n_fatal_attacks_all"] = nkill.gt(0).astype(int)
    events["_n_fatal_attacks_strict"] = (
        events["_n_fatal_attacks_all"] * events["_strict_event"]
    )
    return events


def aggregate_country_month(events: pd.DataFrame) -> pd.DataFrame:
    """Aggregate event-level GTD data to country-month counts."""
    events_with_metrics = add_event_metrics(events)
    metric_map = {
        "n_attacks_all": "_n_attacks_all",
        "n_attacks_strict": "_n_attacks_strict",
        "n_success_all": "_n_success_all",
        "n_success_strict": "_n_success_strict",
        "n_suicide_all": "_n_suicide_all",
        "n_suicide_strict": "_n_suicide_strict",
        "fatalities_all": "_fatalities_all",
        "fatalities_strict": "_fatalities_strict",
        "injuries_all": "_injuries_all",
        "injuries_strict": "_injuries_strict",
        "n_fatal_attacks_all": "_n_fatal_attacks_all",
        "n_fatal_attacks_strict": "_n_fatal_attacks_strict",
    }

    group_cols = ["country", "country_txt", "region", "region_txt", "month"]
    aggregated = (
        events_with_metrics.groupby(group_cols, dropna=False)
        .agg(**{name: (source, "sum") for name, source in metric_map.items()})
        .reset_index()
    )
    return aggregated


def format_values(values: pd.Series) -> str:
    """Format unique metadata values for diagnostics."""
    cleaned = sorted(str(value) for value in values.dropna().unique())
    return "; ".join(cleaned)


def first_mode_or_missing(values: pd.Series) -> object:
    """Return the modal non-missing value, or NA if all values are missing."""
    modes = values.mode(dropna=True)
    if modes.empty:
        return pd.NA
    return modes.iloc[0]


def dataframe_to_markdown(df: pd.DataFrame) -> str:
    """Render a small DataFrame as a markdown table without extra dependencies."""
    if df.empty:
        return ""

    printable = df.copy().astype(str)
    headers = list(printable.columns)
    rows = printable.values.tolist()

    header_line = "| " + " | ".join(headers) + " |"
    separator_line = "| " + " | ".join(["---"] * len(headers)) + " |"
    body_lines = ["| " + " | ".join(row) + " |" for row in rows]
    return "\n".join([header_line, separator_line] + body_lines)


def country_metadata_diagnostics(events: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Create canonical country metadata and flag country code metadata changes."""
    metadata = events[["country", "country_txt", "region", "region_txt"]].drop_duplicates()

    diagnostics = (
        metadata.groupby("country", dropna=False)
        .agg(
            n_country_names=("country_txt", "nunique"),
            country_names=("country_txt", format_values),
            n_regions=("region", "nunique"),
            regions=("region", format_values),
            n_region_names=("region_txt", "nunique"),
            region_names=("region_txt", format_values),
        )
        .reset_index()
    )
    warnings = diagnostics.loc[
        (diagnostics["n_country_names"] > 1)
        | (diagnostics["n_regions"] > 1)
        | (diagnostics["n_region_names"] > 1)
    ].copy()

    canonical = (
        events.sort_values(["country", "month"])
        .groupby("country", dropna=False)
        .agg(
            country_txt=("country_txt", first_mode_or_missing),
            region=("region", first_mode_or_missing),
            region_txt=("region_txt", first_mode_or_missing),
        )
        .reset_index()
    )
    return canonical, warnings


def build_complete_panel(
    aggregated: pd.DataFrame,
    events: pd.DataFrame,
    start: pd.Timestamp,
    end: pd.Timestamp,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Build a complete country-month panel and merge GTD counts onto it."""
    if end < start:
        raise ValueError(f"End month {end.date()} is before start month {start.date()}.")

    canonical_metadata, metadata_warnings = country_metadata_diagnostics(events)
    months = pd.date_range(start=start, end=end, freq="MS")

    countries = canonical_metadata[["country"]].drop_duplicates()
    spine = countries.merge(pd.DataFrame({"month": months}), how="cross")
    spine = spine.merge(canonical_metadata, on="country", how="left")

    count_columns = [
        column
        for column in aggregated.columns
        if column not in ["country", "country_txt", "region", "region_txt", "month"]
    ]

    aggregated_country_month = (
        aggregated.groupby(["country", "month"], dropna=False)[count_columns]
        .sum()
        .reset_index()
    )

    panel = spine.merge(aggregated_country_month, on=["country", "month"], how="left")
    panel[count_columns] = panel[count_columns].fillna(0)

    for column in count_columns:
        panel[column] = panel[column].round().astype("int64")

    panel["attack_month_t"] = panel["n_attacks_strict"].gt(0).astype("int64")

    panel = panel[
        ["country", "country_txt", "region", "region_txt", "month"]
        + count_columns
        + ["attack_month_t"]
    ].sort_values(["country", "month"])
    return panel, metadata_warnings


def save_event_outputs(events: pd.DataFrame, interim_dir: Path) -> None:
    """Save cleaned event-level outputs."""
    events.to_parquet(interim_dir / "gtd_events_clean.parquet", index=False)
    events.head(1000).to_csv(interim_dir / "gtd_events_clean_sample.csv", index=False)


def save_panel_outputs(panel: pd.DataFrame, processed_dir: Path) -> None:
    """Save final country-month panel outputs."""
    panel.to_parquet(processed_dir / "gtd_country_month_panel.parquet", index=False)
    panel.to_csv(processed_dir / "gtd_country_month_panel.csv", index=False)


def make_diagnostic_plots(panel: pd.DataFrame, figures_dir: Path) -> None:
    """Create simple diagnostic plots from the final panel."""
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as exc:
        raise ImportError(
            "matplotlib is required to create diagnostic plots. "
            "Install requirements with: pip install -r requirements.txt"
        ) from exc

    monthly = panel.groupby("month", as_index=False)[
        ["n_attacks_all", "n_attacks_strict"]
    ].sum()

    plt.figure(figsize=(10, 5))
    plt.plot(monthly["month"], monthly["n_attacks_all"], label="All GTD events")
    plt.plot(monthly["month"], monthly["n_attacks_strict"], label="Strict events")
    plt.title("Global GTD Attacks by Month")
    plt.xlabel("Month")
    plt.ylabel("Number of attacks")
    plt.legend()
    plt.tight_layout()
    plt.savefig(figures_dir / "global_attacks_by_month.png", dpi=200)
    plt.close()

    top20 = (
        panel.groupby("country_txt", as_index=False)["n_attacks_strict"]
        .sum()
        .sort_values("n_attacks_strict", ascending=False)
        .head(20)
        .sort_values("n_attacks_strict")
    )

    plt.figure(figsize=(10, 7))
    plt.barh(top20["country_txt"], top20["n_attacks_strict"])
    plt.title("Top 20 Countries by Strict GTD Attacks")
    plt.xlabel("Number of strict attacks")
    plt.ylabel("")
    plt.tight_layout()
    plt.savefig(figures_dir / "top20_countries_attacks_strict.png", dpi=200)
    plt.close()


def diagnostics_markdown(
    events: pd.DataFrame,
    panel: pd.DataFrame,
    diagnostics: dict[str, int],
    missing_columns: list[str],
    metadata_warnings: pd.DataFrame,
    start: pd.Timestamp,
    end: pd.Timestamp,
) -> str:
    """Create a markdown diagnostics report."""
    period_events = events.loc[events["month"].between(start, end)].copy()
    period_events = add_event_metrics(period_events)

    n_countries = panel["country"].nunique(dropna=False)
    n_months = panel["month"].nunique()
    expected_rows = n_countries * n_months
    total_attacks_events = int(period_events["_n_attacks_all"].sum())
    total_attacks_panel = int(panel["n_attacks_all"].sum())
    attack_months = int(panel["attack_month_t"].sum())
    attack_month_share = panel["attack_month_t"].mean()

    top20 = (
        panel.groupby(["country", "country_txt"], as_index=False)["n_attacks_strict"]
        .sum()
        .sort_values("n_attacks_strict", ascending=False)
        .head(20)
    )
    attacks_by_year = (
        period_events.assign(year=period_events["month"].dt.year)
        .groupby("year", as_index=False)
        .agg(
            n_attacks_all=("_n_attacks_all", "sum"),
            n_attacks_strict=("_n_attacks_strict", "sum"),
        )
    )

    optional_missing_notes = []
    for source, metrics in {
        "doubtterr": ["n_attacks_strict"],
        "success": ["n_success_all", "n_success_strict"],
        "suicide": ["n_suicide_all", "n_suicide_strict"],
        "nkill": ["fatalities_all", "fatalities_strict", "n_fatal_attacks_all"],
        "nwound": ["injuries_all", "injuries_strict"],
    }.items():
        if source in missing_columns:
            optional_missing_notes.append(
                f"- `{source}` is missing; related metrics were computed from "
                f"default values for this run: {', '.join(metrics)}."
            )

    lines = [
        "# GTD Country-Month Panel Diagnostics",
        "",
        f"- Selected panel period: {start.date()} to {end.date()}",
        f"- Raw GTD rows loaded: {diagnostics['raw_rows_loaded']:,}",
        (
            "- Rows dropped because `imonth` was missing or 0: "
            f"{diagnostics['dropped_unknown_month']:,}"
        ),
        f"- Rows dropped because `imonth` was outside 1-12: {diagnostics['dropped_invalid_month']:,}",
        f"- Rows dropped because a valid month date could not be created: {diagnostics['dropped_invalid_date']:,}",
        f"- Cleaned event rows: {diagnostics['cleaned_event_rows']:,}",
        f"- Number of countries: {n_countries:,}",
        f"- Number of months: {n_months:,}",
        f"- Expected panel rows = countries x months: {expected_rows:,}",
        f"- Actual panel rows: {len(panel):,}",
        f"- Total attacks in cleaned data within selected period: {total_attacks_events:,}",
        f"- Total attacks in panel after aggregation: {total_attacks_panel:,}",
        f"- Country-months with at least one strict attack: {attack_months:,}",
        f"- Share of country-months with at least one strict attack: {attack_month_share:.3f}",
        "",
        "Note: missing `nkill` and `nwound` values are set to zero only while "
        "aggregating fatalities and injuries. This is a modeling choice to revisit later.",
        "",
        "`attack_month_t` equals 1 when `n_attacks_strict > 0` in the current "
        "country-month. It is not an onset variable.",
        "",
        "## Missing Expected GTD Variables",
        "",
    ]

    if missing_columns:
        lines.extend(f"- `{column}`" for column in missing_columns)
    else:
        lines.append("- None")

    if optional_missing_notes:
        lines.extend(["", "## Optional Variable Warnings", ""])
        lines.extend(optional_missing_notes)

    lines.extend(["", "## Country Metadata Warnings", ""])
    if metadata_warnings.empty:
        lines.append("- No country codes with multiple names or regions detected.")
    else:
        lines.append(
            "The final panel uses the modal country name and region for each country "
            "code. Review these country codes before modeling:"
        )
        lines.append("")
        lines.append(
            dataframe_to_markdown(metadata_warnings)
        )

    lines.extend(["", "## Top 20 Countries by `n_attacks_strict`", ""])
    lines.append(dataframe_to_markdown(top20))

    lines.extend(["", "## Attacks by Year", ""])
    if attacks_by_year.empty:
        lines.append("- No attacks in selected period.")
    else:
        lines.append(dataframe_to_markdown(attacks_by_year))

    return "\n".join(lines) + "\n"


def save_diagnostics(
    events: pd.DataFrame,
    panel: pd.DataFrame,
    diagnostics: dict[str, int],
    missing_columns: list[str],
    metadata_warnings: pd.DataFrame,
    diagnostics_dir: Path,
    start: pd.Timestamp,
    end: pd.Timestamp,
) -> None:
    """Save the diagnostics report."""
    report = diagnostics_markdown(
        events=events,
        panel=panel,
        diagnostics=diagnostics,
        missing_columns=missing_columns,
        metadata_warnings=metadata_warnings,
        start=start,
        end=end,
    )
    (diagnostics_dir / "gtd_panel_diagnostics.md").write_text(report, encoding="utf-8")


def main() -> None:
    """Run the GTD country-month panel build."""
    args = parse_args()
    start = month_start(args.start)
    end = month_start(args.end)

    ensure_output_dirs(
        [
            args.raw_dir,
            args.interim_dir,
            args.processed_dir,
            args.diagnostics_dir,
            args.figures_dir,
        ]
    )

    print(
        textwrap.dedent(
            f"""
            Building GTD country-month panel
            Period: {start.date()} to {end.date()}
            """
        ).strip()
    )

    raw_file = find_gtd_file(args.raw_dir)
    raw_df = load_gtd_file(raw_file)
    available_columns, missing_columns = validate_columns(raw_df)

    events, diagnostics = clean_event_data(raw_df, available_columns)
    print(f"Cleaned event rows: {len(events):,}")
    save_event_outputs(events, args.interim_dir)

    aggregated = aggregate_country_month(events)
    panel, metadata_warnings = build_complete_panel(aggregated, events, start, end)
    save_panel_outputs(panel, args.processed_dir)
    save_diagnostics(
        events=events,
        panel=panel,
        diagnostics=diagnostics,
        missing_columns=missing_columns,
        metadata_warnings=metadata_warnings,
        diagnostics_dir=args.diagnostics_dir,
        start=start,
        end=end,
    )
    make_diagnostic_plots(panel, args.figures_dir)

    print(f"Final panel rows: {len(panel):,}")
    print(f"Saved panel parquet: {args.processed_dir / 'gtd_country_month_panel.parquet'}")
    print(f"Saved panel CSV: {args.processed_dir / 'gtd_country_month_panel.csv'}")
    print(f"Saved diagnostics: {args.diagnostics_dir / 'gtd_panel_diagnostics.md'}")


if __name__ == "__main__":
    main()
